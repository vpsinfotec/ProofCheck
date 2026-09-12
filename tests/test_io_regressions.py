from concurrent.futures import ThreadPoolExecutor
import json
import time

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from proofcheck import excel, images, ocr, ocr_cache, report_xlsx
from proofcheck.models import ColumnResult, MatchResult, Meta, RunResult, Status, Summary


def workbook(tmp_path, rows):
    path = tmp_path / 'input.xlsx'
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(path)
    return str(path)


def test_duplicate_headers_rejected(tmp_path):
    path = workbook(tmp_path, [['Name', 'Name'], ['Alice', 'Bob']])
    with pytest.raises(excel.ExcelError, match='Duplicate'):
        excel.load_columns(path, all_columns=True)


def test_duplicate_selection_deduplicated_and_header_row_honored(tmp_path):
    path = workbook(tmp_path, [['Title'], ['Name'], ['Alice']])
    assert excel.inspect(path, header_row=2)['Sheet'] == ['Name']
    cols = excel.load_columns(path, columns=['Name', 'Name'], header_row=2)
    assert len(cols) == 1 and cols[0].cells == [(3, 'Alice')]


def test_expanded_workbook_limit(tmp_path, monkeypatch):
    path = workbook(tmp_path, [['Name'], ['Alice']])
    monkeypatch.setattr(excel, 'MAX_WORKBOOK_BYTES', 10)
    with pytest.raises(excel.ExcelError, match='Expanded'):
        excel.load_columns(path, all_columns=True)


def test_selected_cell_limit(tmp_path, monkeypatch):
    path = workbook(tmp_path, [['Name'], ['Alice'], ['Bob']])
    monkeypatch.setattr(excel, 'MAX_CELLS', 1)
    with pytest.raises(excel.ExcelError, match='limit'):
        excel.load_columns(path, all_columns=True)


def test_report_formula_text_and_safe_unique_titles(tmp_path):
    result = RunResult(Meta('=1+1', 'file.pdf', Meta.now_iso(), 90, {}), Summary(), [
        ColumnResult(name, [MatchResult(2, '=HYPERLINK("https://invalid.test")', Status.MISSING)])
        for name in ['A/B', 'A?B', 'Summary', 'x' * 32, 'x' * 33]
    ])
    path = tmp_path / 'report.xlsx'
    report_xlsx.write(result, str(path))
    wb = load_workbook(path)
    assert len(set(n.casefold() for n in wb.sheetnames)) == 6
    for ws in list(wb)[1:]:
        assert ws['B2'].data_type == 's' and ws['B2'].value.startswith('=')
        assert len(ws.title) <= 31
    wb.close()


@pytest.mark.parametrize('content', ['[]', '{"x": "text"}', '{"1": 3}', '{"0":"bad"}', '{'])
def test_corrupt_cache_is_a_miss(content):
    directory = ocr_cache.cache_dir()
    directory.mkdir()
    path = ocr_cache._entry_path(directory, 'a' * 64, 300, 'eng', 6)
    path.write_text(content)
    assert ocr_cache.load('a' * 64, dpi=300, lang='eng', psm=6) is None


def test_cache_atomic_under_concurrent_writes():
    def write(i):
        ocr_cache.store('a' * 64, dpi=300, lang='eng', pages={1: str(i)}, psm=6)
    with ThreadPoolExecutor(max_workers=8) as pool:
        list(pool.map(write, range(80)))
    result = ocr_cache.load('a' * 64, dpi=300, lang='eng', psm=6)
    assert result is not None and result[1].isdigit()
    assert not list(ocr_cache.cache_dir().glob('*.tmp'))


def test_image_ocr_failure_not_cached(tmp_path, monkeypatch):
    path = tmp_path / 'image.png'
    Image.new('RGB', (10, 10), 'white').save(path)
    monkeypatch.setattr(ocr, 'available', lambda: True)
    def fail(*a, **kw):
        raise ocr.OcrError('temporary failure')
    monkeypatch.setattr(ocr, 'ocr_image_file', fail)
    assert images.extract(str(path)).ocr_error
    monkeypatch.setattr(ocr, 'ocr_image_file', lambda *a, **kw: 'Recovered')
    assert images.extract(str(path)).pages[1] == 'Recovered'
    monkeypatch.setattr(ocr, 'available', lambda: False)
    assert images.extract(str(path)).pages[1] == 'Recovered'


def test_uniform_page_skips_tesseract(monkeypatch):
    def fail(*a, **kw):
        raise AssertionError('Blank image should not invoke Tesseract')
    monkeypatch.setattr(ocr._pytesseract, 'image_to_data', fail)
    image = Image.new('L', (40, 40), 255)
    assert ocr._best_ocr(image, lang='eng', psm=6, oem=3)[0] == ''


def test_multiframe_image_explicitly_rejected(tmp_path):
    path = tmp_path / 'frames.tiff'
    Image.new('RGB', (10, 10)).save(path, save_all=True, append_images=[Image.new('RGB', (10, 10), 'red')])
    with pytest.raises(ocr.OcrError, match='Multi-frame'):
        ocr._load_image_file(str(path))


def test_oversized_image_rejected_before_decode(tmp_path, monkeypatch):
    path = tmp_path / 'large.png'
    Image.new('RGB', (20, 20)).save(path)
    monkeypatch.setattr(ocr, 'MAX_IMAGE_PIXELS', 100)
    with pytest.raises(ocr.OcrError, match='pixels'):
        ocr._load_image_file(str(path))
