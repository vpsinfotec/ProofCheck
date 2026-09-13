import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas

from proofcheck.matcher import PreparedMatcher, match_value
from proofcheck.models import PageOccurrences, RepeatedWord, Status
from proofcheck.web.app import app
from proofcheck.web.schemas import SummaryModel


@pytest.mark.parametrize("text,word", [
    ("Areeb Areeb Khan", "areeb"),
    ("Areeb Khan Khan", "khan"),
])
def test_repeated_first_or_last_word_is_flagged_even_when_exact(text, word):
    r = match_value("Areeb Khan", {1: text})
    assert r.status is Status.EXACT and r.score == 100
    assert r.needs_review and r.occurrence_count == 1
    assert r.repeated_words == [RepeatedWord(word, 2, page=1)]


def test_repeated_middle_word_is_reviewed_without_silently_correcting_verdict():
    r = match_value("Areeb Ahmed Khan", {3: "Areeb Ahmed Ahmed Khan"}, fuzzy_threshold=100)
    assert r.status is Status.MISSING  # review collapse must not produce an exact match
    assert r.needs_review and r.occurrence_count == 0
    assert r.repeated_words == [RepeatedWord("ahmed", 2, page=3)]


@pytest.mark.parametrize("pages,counts", [
    ({1: "Areeb Khan\nAreeb Khan"}, [(1, 2)]),
    ({9: "Areeb Khan", 2: "Areeb Khan"}, [(2, 1), (9, 1)]),
    ({1: "Areeb Khan; Areeb Khan", 2: "Areeb Khan"}, [(1, 2), (2, 1)]),
])
def test_repeated_full_name_on_same_or_different_pages(pages, counts):
    r = match_value("Areeb Khan", pages)
    assert r.page == counts[0][0] and r.status is Status.EXACT
    assert r.occurrences == [PageOccurrences(page, count) for page, count in counts]
    assert r.occurrence_count == sum(count for _, count in counts)
    assert r.needs_review and not r.repeated_words


@pytest.mark.parametrize("value,word", [("Areeb Areeb Khan", "areeb"), ("Areeb Khan Khan", "khan")])
def test_repetition_in_spreadsheet_is_also_reviewed(value, word):
    r = match_value(value, {1: "Areeb Khan"})
    assert r.repeated_words == [RepeatedWord(word, 2)]
    assert r.needs_review


def test_later_page_repeated_word_is_found_after_earlier_exact():
    r = match_value("Areeb Khan", {1: "Areeb Khan", 4: "Areeb Areeb Areeb Khan Khan"})
    assert r.page == 1
    assert r.repeated_words == [RepeatedWord("areeb", 3, 4), RepeatedWord("khan", 2, 4)]


def test_counts_do_not_conflate_names_with_longer_word_prefixes():
    r = match_value("Areeb Khan", {1: "Areeb Khan; Areeb Khanna; Shareeb Khan"})
    assert r.occurrence_count == 1 and not r.needs_review
    assert match_value("Ann", {1: "Ann, Joanna and Annette"}).occurrence_count == 1


def test_unrelated_repetition_does_not_flag_name():
    r = match_value("Areeb Khan", {1: "Other Other person; Areeb Khan; Khanam Khanam"})
    assert not r.repeated_words and not r.needs_review


def test_normalization_and_ocr_whitespace_apply_to_duplicate_checks():
    r = match_value("Aréeb Khan", {1: "AREEB\tAREEB\nKHAN\nAreeb—Khan"},
                    fold_diacritics=True, strip_punctuation=True)
    assert r.occurrence_count == 2
    assert r.repeated_words == [RepeatedWord("areeb", 2, 1)]


def test_hyphenated_words_and_unicode_repetitions():
    r = match_value("Anne-Marie Khan", {1: "Anne-Marie Anne-Marie Khan"})
    assert r.repeated_words == [RepeatedWord("anne-marie", 2, 1)]
    r = match_value("عريب خان", {1: "عريب عريب خان"})
    assert r.repeated_words == [RepeatedWord("عريب", 2, 1)]


def test_reverse_counts_do_not_double_count_overlapping_or_palindromic_names():
    r = match_value("Areeb Khan", {1: "Areeb Khan Areeb", 2: "Khan Areeb"}, reverse=True)
    assert r.occurrences == [PageOccurrences(1, 1), PageOccurrences(2, 1)]
    assert match_value("Areeb Areeb", {1: "Areeb Areeb"}, reverse=True).occurrence_count == 1


def test_reverse_duplicate_word_detection():
    r = match_value("Areeb Khan", {1: "Khan Khan Areeb"}, reverse=True)
    assert r.repeated_words == [RepeatedWord("khan", 2, 1)]


def test_duplicate_review_metadata_is_reused_but_lists_are_isolated(monkeypatch):
    m = PreparedMatcher({1: "Areeb Areeb Khan; Areeb Khan"})
    a = m.match("Areeb Khan", row=2)
    monkeypatch.setattr(m.duplicate_auditor, "inspect", lambda *a: pytest.fail("Must reuse review"))
    a.occurrences.clear()
    a.repeated_words.clear()
    b = m.match("Areeb Khan", row=3)
    assert b.row == 3 and b.occurrence_count == 2 and b.repeated_words and b.needs_review


@pytest.mark.parametrize("value,pages", [(None, {1: "Areeb Areeb Khan"}), ("", {}), ("Areeb Khan", {})])
def test_empty_values_and_pages_never_report_duplicates(value, pages):
    r = match_value(value, pages)
    assert r.occurrence_count == 0 and not r.needs_review


def test_legacy_history_is_not_falsely_reported_as_audited():
    assert SummaryModel(total=1, exact=1, fuzzy=0, missing=0, skipped=0, pass_rate=1).duplicate_review is None


def test_real_upload_duplicate_findings_survive_api_history_and_reports():
    excel, pdf = io.BytesIO(), io.BytesIO()
    wb = Workbook()
    wb.active.append(["Name"])
    for name in ("Areeb Khan", "Areeb Areeb Khan", "Areeb Khan Khan", "Sana Ali"):
        wb.active.append([name])
    wb.save(excel)
    c = canvas.Canvas(pdf)
    for lines in (["Areeb Areeb Khan", "Areeb Khan Khan", "Sana Ali"], ["Areeb Khan"]):
        for i, line in enumerate(lines):
            c.drawString(40, 800 - i * 30, line)
        c.showPage()
    c.save()
    with TestClient(app) as client:
        response = client.post("/api/check", files={
            "excel": ("names.xlsx", excel.getvalue()), "pdf": ("names.pdf", pdf.getvalue())
        }, data={"columns_json": '["Name"]'})
        assert response.status_code == 200, response.text
        data = response.json()
        assert data["summary"]["duplicate_review"] == 3
        r = data["columns"][0]["results"][0]
        assert r["needs_review"] and r["occurrence_count"] == 3
        assert r["occurrences"] == [{"page": 1, "count": 2}, {"page": 2, "count": 1}]
        assert {w["word"] for w in r["repeated_words"]} == {"areeb", "khan"}
        html = client.get(data["report_urls"]["html"]).text
        assert "Review duplicates" in html and "page 1: 2; page 2: 1" in html
        assert "Review repeated word in spreadsheet value" in html
        report = load_workbook(io.BytesIO(client.get(data["report_urls"]["xlsx"]).content))
        assert "Review duplicates" in report["Name"]["C2"].value
        assert "page 1: 2; page 2: 1" in report["Name"]["E2"].value
        assert "Review repeated word in spreadsheet value" in report["Name"]["E3"].value
        history = client.get("/api/history").json()
        assert history["runs"][0]["summary"]["duplicate_review"] == 3
